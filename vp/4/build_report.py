"""Build the Lab 4 report only from saved real measurements and screenshots."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
data = json.loads((ROOT / 'results/summary.json').read_text())
rows = data['rows']
assert {(r['replicas'], r['target_rps']) for r in rows} >= {
    (n, rate) for n in (1, 2, 3) for rate in (10, 25, 50, 100)
}, 'All required real tests must finish before building the report'
for row in rows:
    stem = f"r{row['replicas']}-rate{row['target_rps']}"
    assert (ROOT / 'results' / f'{stem}.json').exists()
    assert (ROOT / 'results' / f'{stem}.prometheus.json').exists()


def accepted(row):
    return (row['p95_ms'] < 500 and row['errors_percent'] < 1
            and row['dropped_iterations'] == 0
            and row['actual_rps'] >= 0.95 * row['target_rps'])


best = {}
for replicas in (1, 2, 3):
    candidates = [r for r in rows if r['replicas'] == replicas and accepted(r)]
    best[replicas] = max(candidates, key=lambda r: r['actual_rps']) if candidates else None

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.left_margin, section.right_margin = Cm(3), Cm(1.5)
section.top_margin, section.bottom_margin = Cm(2), Cm(1.3)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.space_after = Pt(5)
for name in ('Title', 'Heading 1'):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.font.size = Pt(16 if name == 'Title' else 14)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)
title_ppr = doc.styles['Title']._element.get_or_add_pPr()
border = title_ppr.find(qn('w:pBdr'))
if border is not None:
    title_ppr.remove(border)


def paragraph(value):
    return doc.add_paragraph(value)


def heading(value):
    doc.add_paragraph(value, style='Heading 1')


def caption(value, above=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6 if above else 2)
    p.paragraph_format.space_after = Pt(2 if above else 8)
    run = p.add_run(value)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)


figure = 0


def picture(name, description, width=15.5):
    global figure
    path = ROOT / 'screenshots' / name
    assert path.exists(), path
    figure += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    caption(f'Рисунок {figure}. {description}')


table_number = 0


def table(description, headers, values, widths=None, font_size=9):
    global table_number
    table_number += 1
    caption(f'Таблица {table_number}. {description}', above=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = str(value)
    for row in values:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for tr in t.rows:
        for i, cell in enumerate(tr.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(font_size)


def fmt(value, digits=2):
    return f'{value:.{digits}f}' if value is not None else 'нет данных'


title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Лабораторная работа № 4. Горизонтальное масштабирование веб-приложения')
paragraph('Цель работы: проверить, как изменение числа реплик auth-service влияет на пропускную способность, задержку и использование ресурсов при одинаковой нагрузке.')

heading('1. Подготовка окружения')
paragraph('Я использовал прежний локальный K3s, k6, Prometheus и Grafana. Для ЛР4 развернул отдельную копию auth-service в namespace vp-lr4. Исходный проект: https://github.com/Wladisl4W/Y3-S2-Kubernetes. В приложении добавил GET /work?iterations=20000: он 20000 раз вычисляет SHA-256 и возвращает имя Pod. Во всех тестах число итераций одинаковое. Метрики /metrics сохранились.')
paragraph('Приложение доступно через NodePort 30083. k6 отправлял запросы на адрес узла, так как localhost:30083 на этой машине не отвечал. Для каждой конфигурации я дождался готовности Deployment и проверил работающие Pod. Вывод команд ниже снят в настоящем терминале Ubuntu; команды просмотра сохранённых результатов обозначены прямо.')
picture('cluster-pods.png', 'Работающие Pod приложения после испытаний')
for n in (1, 2, 3):
    picture(f'r{n}-pods.png', f'Сохранённый список Pod при {n} репликах')

heading('2. Сценарий нагрузочного тестирования')
paragraph('Я использовал один файл test-scaling.js для 1, 2 и 3 реплик. Режим constant-arrival-rate задавал 10, 25, 50 и 100 запросов/с на 60 секунд. Один запуск итерации содержал один GET /work и проверку HTTP 200. Между прогонами была пауза 15 секунд. k6 ограничен 400 VU. Для определения устойчивого уровня я проверял p95 < 500 мс, ошибок < 1%, отсутствие dropped_iterations и фактический RPS не ниже 95% целевого.')
paragraph('Исходные выводы и JSON k6, а также ответы Prometheus сохранены в папке results. На следующих снимках видна команда cat, которая открывает сохранённый вывод завершённого прогона при 100 запросах/с.')
for n in (1, 2, 3):
    picture(f'r{n}-rate100.png', f'Просмотр результата k6 для {n} реплик и цели 100 запросов/с')

heading('3. Результаты для одной и нескольких реплик')
table('Измерения с одной репликой',
      ['Цель, RPS', 'Факт, RPS', 'p95, мс', 'Ошибки, %', 'CPU, ядра', 'Dropped'],
      [(r['target_rps'], fmt(r['actual_rps']), fmt(r['p95_ms']),
        fmt(r['errors_percent']), fmt(r['prometheus_cpu_cores'], 3),
        r['dropped_iterations'])
       for r in rows if r['replicas'] == 1],
      [2.1, 2.3, 2.2, 2.4, 2.7, 2.0])
table('Все конфигурации при одинаковой нагрузке',
      ['Реплики', 'Цель, RPS', 'Факт, RPS', 'p95, мс', 'Ошибки, %', 'CPU, ядра', 'Память, МиБ'],
      [(r['replicas'], r['target_rps'], fmt(r['actual_rps']), fmt(r['p95_ms']),
        fmt(r['errors_percent']), fmt(r['prometheus_cpu_cores'], 3),
        fmt(r['prometheus_memory_mib'], 1))
       for r in sorted(rows, key=lambda x: (x['replicas'], x['target_rps']))],
      [1.6, 2.0, 2.1, 2.1, 2.0, 2.7, 3.0], 8)
table('Максимальные измеренные устойчивые уровни',
      ['Реплики', 'Устойчивый RPS', 'p95, мс', 'Ошибки, %'],
      [(n, fmt(best[n]['actual_rps']) if best[n] else 'не найден',
        fmt(best[n]['p95_ms']) if best[n] else 'не найден',
        fmt(best[n]['errors_percent']) if best[n] else 'не найден')
       for n in (1, 2, 3)], [2.0, 4.2, 4.0, 4.0])

heading('4. Анализ в Grafana')
paragraph('Я перенёс dashboard High-load Application из ЛР3 на метки vp-lr4 и маршрут work. Он показывает RPS, p95, ошибки, CPU, память и состояние Pod. Добавил панели доступных реплик, RPS по Pod и CPU по Pod. Фильтры namespace и pod оставил. Источник данных Prometheus после перезапуска Grafana проверен успешным запросом API.')
for n in (1, 2, 3):
    picture(f'grafana-r{n}.png', f'Графики Grafana за сохранённый период испытания с {n} репликами')
picture('grafana-r3-resources.png', 'CPU, память и состояние Pod при трёх репликах')
picture('grafana-r3-replicas.png', 'Доступность трёх реплик и распределение запросов между Pod')
paragraph('CPU и память в таблице взяты из Prometheus по контейнерам auth-service. Для сравнения использован средний уровень за последние 30 секунд соответствующего прогона. Кратковременный всплеск и простой между тестами на графиках поэтому не подменяют эти числа.')
paragraph('При цели 50 запросов/с p95 k6 уменьшился с 2170 мс при одной реплике до 154 мс при двух и 72 мс при трёх. Суммарный CPU при этом составил примерно 0,92, 1,23 и 1,07 ядра соответственно. Память росла вместе с числом Pod: около 30, 60 и 84 МиБ. Это показывает выигрыш в задержке ценой дополнительных ресурсов. Длительность каждого этапа составляла одну минуту, поэтому вывод об устойчивости относится только к этому интервалу.')
if 'connect: connection refused' in (ROOT / 'results/r1-rate100.txt').read_text():
    paragraph('На высоком уровне с одной репликой k6 записал ошибки «connect: connection refused». В сохранённом файле results/r1-pod-describe.txt видны три неудачные liveness probe и последующий перезапуск контейнера. CPU приблизился к одному ядру. Поэтому доля ошибок k6 здесь не равна панели Error Rate: та панель считает только HTTP 5xx, а отказ соединения не содержит ответа HTTP. Перезапуск и недоступность Pod подтверждены; влияние сетевого пути NodePort отдельно не измерялось.')

heading('5. Ускорение и эффективность масштабирования')
if all(best.values()):
    base = best[1]['actual_rps']
    table('Эффективность по максимальным измеренным устойчивым уровням',
          ['Реплики', 'Устойчивый RPS', 'S(N)', 'E(N), %'],
          [(n, fmt(best[n]['actual_rps']), fmt(best[n]['actual_rps'] / base),
            fmt(best[n]['actual_rps'] / base / n * 100)) for n in (1, 2, 3)],
          [2.0, 4.8, 3.0, 3.7])
    paragraph('Я вычислил S(N) = RPSmax(N) / RPSmax(1), E(N) = S(N) / N × 100%. Эти значения относятся только к проверенным уровням нагрузки и одинаковым порогам p95, ошибок и dropped_iterations.')
else:
    paragraph('Для одной из конфигураций устойчивый уровень в проверенном диапазоне не найден. Поэтому численный коэффициент ускорения для неё не определялся; придумывать значение нельзя.')

max_tested = max(r['target_rps'] for r in rows)
if all(best.values()) and all(best[n]['target_rps'] == max_tested for n in (1, 2, 3)):
    paragraph('Наибольший проверенный уровень прошли все три конфигурации. Поэтому указанные максимумы являются только нижними оценками предельной пропускной способности. Они не доказывают отсутствие пользы от масштабирования: тест не довёл систему до общего предела при одинаковых критериях.')
else:
    paragraph('Сравнение учитывает те уровни, которые выдержали одинаковые критерии. Поскольку все Pod работают на одном узле K3s и делят его CPU и память, эффект масштабирования может ограничиваться ресурсами узла. Вывод о конкретном узком месте делаю только там, где он подтверждается метриками.')
    paragraph('На измеренных уровнях рост от одной до двух реплик дал приблизительно двукратный устойчивый RPS, а третья не подняла максимальный подтверждённый уровень выше 50 RPS. Это не доказывает, что третья бесполезна: промежуточные уровни между 50 и 100 RPS не проверялись. При 100 RPS три реплики дали 98,04 фактических RPS без HTTP-ошибок, но p95 достиг 1200 мс и было 94 пропущенных итерации. Дальнейшее добавление реплик без нового опыта на более мелких шагах и оценки ресурсов общего узла я не могу признать обоснованным.')

heading('6. Ответы на контрольные вопросы')
qa = [
    ('1. Чем горизонтальное масштабирование отличается от вертикального?',
     'Горизонтальное добавляет экземпляры приложения, вертикальное увеличивает ресурсы одного экземпляра.'),
    ('2. Почему число Pod не гарантирует пропорционального роста RPS?',
     'Pod могут делить один CPU, базу данных, сеть и точку входа. Генератор нагрузки тоже может стать ограничением.'),
    ('3. Почему нужны одинаковые условия тестов?',
     'Иначе изменение RPS и задержки нельзя связать именно с количеством реплик.'),
    ('4. Что показывают ускорение и эффективность?',
     'Ускорение сравнивает предельный устойчивый RPS с одной репликой, эффективность показывает этот выигрыш на одну реплику.'),
    ('5. Как найти предел устойчивой производительности?',
     'Повышать целевую интенсивность и проверять фактический RPS, p95, ошибки, dropped_iterations и динамику задержки в течение этапа.'),
    ('6. Почему масштабирование бывает неэффективным при исправном Kubernetes?',
     'Дополнительные Pod не устраняют общий дефицит CPU, блокировки, базу данных или ограничение генератора.'),
]
for question, answer in qa:
    p = paragraph(question)
    p.paragraph_format.keep_with_next = True
    paragraph(answer)

heading('Вывод')
if all(best.values()):
    paragraph('Я проверил масштабирование auth-service от одной до трёх реплик одинаковым сценарием k6. По сохранённым измерениям наибольшие устойчивые уровни составили '
              + ', '.join(f"{n} репл.: {best[n]['actual_rps']:.2f} RPS" for n in (1, 2, 3))
              + '. При оценке дальнейшего масштабирования нужно учитывать, достигнут ли общий предел нагрузки и не ограничивает ли результат один узел K3s.')
else:
    paragraph('Я проверил одинаковые уровни нагрузки при одной, двух и трёх репликах. Для конфигураций без устойчивого уровня в выбранном диапазоне нужен дополнительный опыт с другой интенсивностью. Данные и скриншоты сохранены без реконструкции.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Cm(0)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.save(ROOT / 'ЛР4_отчёт.docx')
