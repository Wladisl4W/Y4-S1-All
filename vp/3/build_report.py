"""Build the Lab 3 report from genuine experiment files and screenshots."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


folder = Path(__file__).parent
rows = json.loads((folder / 'results/summary.json').read_text())['rows']
assert [r['vus'] for r in rows] == [10, 25, 50, 100, 200]
by_vu = {r['vus']: r for r in rows}
for r in rows:
    assert (folder / 'results' / f"vu{r['vus']}.json").exists()
    assert (folder / 'results' / f"vu{r['vus']}.prometheus.json").exists()

dashboard = json.loads((folder / 'high-load-dashboard.json').read_text())
assert len(dashboard['panels']) == 7

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.left_margin, section.right_margin = Cm(3), Cm(1.5)
section.top_margin, section.bottom_margin = Cm(2), Cm(1.3)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.keep_together = True
for name in ('Title', 'Heading 1'):
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.font.size = Pt(16 if name == 'Title' else 14)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)


def paragraph(text):
    return doc.add_paragraph(text)


def heading(text):
    doc.add_paragraph(text, style='Heading 1')


def caption(text, above=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6 if above else 2)
    p.paragraph_format.space_after = Pt(2 if above else 8)
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)


figure_number = 0


def picture(name, description, width=15.5):
    global figure_number
    path = folder / 'screenshots' / name
    assert path.exists(), path
    figure_number += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    caption(f'Рисунок {figure_number}. {description}')


table_number = 0


def table(description, headers, items, font_size=9, widths=None):
    global table_number
    table_number += 1
    caption(f'Таблица {table_number}. {description}', above=True)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.autofit = False
    if widths:
        for column, width in zip(t.columns, widths):
            column.width = Cm(width)
    for i, head in enumerate(headers):
        t.rows[0].cells[i].text = str(head)
    for item in items:
        cells = t.add_row().cells
        for i, value in enumerate(item):
            cells[i].text = str(value)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.space_after = Pt(1)
                p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0 or description.startswith('Запросы')
                               else WD_ALIGN_PARAGRAPH.CENTER)
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(font_size)
    return t


def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Liberation Mono'
    run.font.size = Pt(9)


title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Лабораторная работа № 3. Визуализация и анализ нагрузки в Grafana')
title_ppr = doc.styles['Title']._element.get_or_add_pPr()
border = title_ppr.find(qn('w:pBdr'))
if border is not None:
    title_ppr.remove(border)
paragraph('Цель работы: показать метрики приложения в Grafana, сравнить их с результатами k6 и проверить, при какой нагрузке заметно ухудшается время ответа.')

heading('1. Установка и проверка Grafana')
paragraph('Я продолжил работу с auth-service из ЛР1 и ЛР2. Исходный проект: https://github.com/Wladisl4W/Y3-S2-Kubernetes. Для этой лабораторной я развернул отдельную копию приложения в namespace vp-lr3 на прежнем локальном K3s. Grafana включена в kube-prometheus-stack в namespace monitoring. ServiceMonitor собирает метрики auth-service каждые 5 секунд.')
paragraph('Команды, которыми я проверил компоненты и открыл интерфейс:')
code('KUBECONFIG=$HOME/.kube/config k3s kubectl get pods -n monitoring')
code('KUBECONFIG=$HOME/.kube/config k3s kubectl get svc -n monitoring prometheus-grafana')
code('KUBECONFIG=$HOME/.kube/config k3s kubectl -n monitoring port-forward svc/prometheus-grafana 3000:80')
paragraph('После port-forward я открыл http://127.0.0.1:3000 и вошёл под admin. Пароль получил из Kubernetes Secret во время входа. В отчёт и снимки он не попал. Команды и вывод ниже сняты в настоящем терминале Ubuntu из каталога vp/3.')
picture('monitoring-pods.png', 'Под Grafana и компоненты мониторинга работают в namespace monitoring')
picture('grafana-service.png', 'Сервис Grafana в namespace monitoring')
picture('grafana-home.jpg', 'Главная страница Grafana после входа')
paragraph('В Connections → Data sources автоматически появился Prometheus. Кнопка Test вернула сообщение «Successfully queried the Prometheus API».')
picture('prometheus-datasource-test.jpg', 'Успешная проверка источника данных Prometheus')

heading('2. Dashboard High-load Application')
paragraph('Я создал dashboard High-load Application с семью панелями. Переменная namespace выбирает пространство имён, pod выбирает под; в запросах используются реальные метки auth-service. Выбран интервал Last 15 minutes. Для p95 взято окно 1 минута, потому что каждый уровень нагрузки длится 90 секунд.')
panel_descriptions = [
    ('RPS', 'Time series', 'req/s', 'Сколько HTTP-запросов к главной странице обрабатывается за секунду.'),
    ('HTTP 5xx Errors', 'Time series', 'req/s', 'Скорость ответов сервера с кодом 5xx.'),
    ('Error Rate', 'Time series', '%', 'Доля 5xx среди запросов; задан порог 1%.'),
    ('p95 Latency', 'Time series', 'с', 'Оценка времени обработки 95% запросов по гистограмме приложения.'),
    ('CPU Usage', 'Time series', 'ядра', 'Использование CPU контейнером auth-service.'),
    ('Memory Usage', 'Time series', 'МиБ', 'Рабочая память контейнера auth-service.'),
    ('Pod Status', 'Stat', '0/1', 'Доступность цели: 1 означает UP, 0 означает DOWN.'),
]
table('Панели dashboard', ['Панель', 'Тип', 'Ед.', 'Что показывает'], panel_descriptions, 9)
queries = [(p['title'], p['targets'][0]['expr']) for p in dashboard['panels']]
table('Запросы PromQL из экспортированного dashboard', ['Панель', 'PromQL'], queries, 8,
      [3.5, 12.0])
paragraph('Значения CPU и памяти относятся именно к контейнеру auth-service. График Pod Status показывает состояние цели Prometheus с метками namespace и pod. Для чисел 0 и 1 настроены подписи DOWN и UP.')
picture('grafana-100vu-top.jpg', 'RPS, 5xx, доля ошибок и p95 во время этапа 100 VU')
picture('grafana-100vu-resources.jpg', 'CPU, память и Pod Status во время этапа 100 VU')

heading('3. Нагрузочный эксперимент и сравнение k6 с Grafana')
paragraph('Я провёл пять последовательных прогонов k6 с постоянной нагрузкой 10, 25, 50, 100 и 200 VU. Каждый длился 90 секунд, между ними была пауза 15 секунд. Сценарий делает GET /, проверяет код 200 и ждёт одну секунду. Запуск выполнялся из каталога vp/3 командой python3 run_tests.py; этот скрипт для каждого уровня вызывал k6 run --summary-export results/vuN.json test-stages.js и сохранял ответы Prometheus. Исходные файлы лежат в results.')
table('Результаты k6 и Grafana', ['VU', 'RPS k6', 'p95 k6, мс', 'RPS Grafana', 'p95 Grafana, мс', 'CPU, ядра', 'Память, МиБ'], [
    (r['vus'], f"{r['rps_k6']:.2f}", f"{r['p95_k6_ms']:.2f}",
     f"{r['rps_grafana']:.2f}", f"{r['p95_grafana_ms']:.2f}",
     f"{r['cpu_cores']:.3f}", f"{r['memory_mib']:.1f}") for r in rows
], 8)
paragraph('RPS и p95 k6 взяты из итогового JSON каждого прогона. Значения Grafana рассчитаны теми же PromQL-запросами, что стоят в панелях, и усреднены по сохранённым точкам последних 30 секунд этапа. Это важно: k6 считает за весь прогон, а графики берут минутное окно. Поэтому числа близки, но не обязаны совпадать.')
paragraph('На снимках терминала я открыл сохранённые фрагменты вывода завершённых прогонов. Видны текущий каталог, введённая команда cat и реальные строки k6. Это не изображение повторного запуска тестов.')
for vus in (10, 25, 50, 100, 200):
    picture(f'vu{vus}.png', f'Просмотр сохранённого вывода k6 для {vus} VU')
hundred = by_vu[100]
diff = hundred['p95_k6_ms'] - hundred['p95_grafana_ms']
percent = diff / hundred['p95_k6_ms'] * 100
paragraph(f"При 100 VU p95 k6 составил {hundred['p95_k6_ms']:.2f} мс, а p95 Grafana {hundred['p95_grafana_ms']:.2f} мс. Разница {diff:.2f} мс, или {percent:.1f}% от результата k6. k6 измеряет полное время ответа со стороны клиента, включая сетевой путь и ожидание до обработки. Гистограмма приложения измеряет только обработку внутри auth-service. Дополнительно методы вычисления квантиля и интервалы отличаются. По этим данным нельзя считать одно значение ошибочным.")

heading('4. Анализ точки ухудшения')
paragraph('CPU начал расти уже с 10 VU и продолжал расти почти на каждом шаге. Память увеличивалась медленнее. RPS рос на всех этапах; переход от 100 к 200 VU снова почти удвоил его. Ответов 5xx и ошибок k6 не было. Поэтому точку насыщения сервера в этом опыте я не обнаружил.')
paragraph('Рост p95 k6 впервые заметен на 50 VU, а сильный скачок произошёл на 100 VU. На графике я отметил начало этапа 100 VU аннотацией Grafana. Это отметка ухудшения клиентского времени ответа, а не доказанная граница пропускной способности сервера.')
before, after = by_vu[50], by_vu[100]
doc.add_page_break()
table('До и после заметного роста клиентского p95', ['Показатель', '50 VU', '100 VU'], [
    ('VU', '50', '100'),
    ('RPS k6', f"{before['rps_k6']:.2f}", f"{after['rps_k6']:.2f}"),
    ('p95 k6, мс', f"{before['p95_k6_ms']:.2f}", f"{after['p95_k6_ms']:.2f}"),
    ('p99 k6, мс', f"{before['p99_k6_ms']:.2f}", f"{after['p99_k6_ms']:.2f}"),
    ('Ошибки k6, %', f"{before['errors_k6_percent']:.0f}", f"{after['errors_k6_percent']:.0f}"),
    ('CPU, ядра', f"{before['cpu_cores']:.3f}", f"{after['cpu_cores']:.3f}"),
    ('Память, МиБ', f"{before['memory_mib']:.1f}", f"{after['memory_mib']:.1f}"),
], 10)
picture('degradation-annotation-top.jpg', 'Аннотация Grafana отмечает начало этапа 100 VU')
picture('degradation-annotation-resources.jpg', 'CPU и память на графике с той же отметкой')
paragraph(f"При переходе с 50 на 100 VU p95 k6 вырос с {before['p95_k6_ms']:.2f} до {after['p95_k6_ms']:.2f} мс, CPU с {before['cpu_cores']:.3f} до {after['cpu_cores']:.3f} ядра, память с {before['memory_mib']:.1f} до {after['memory_mib']:.1f} МиБ. RPS вырос с {before['rps_k6']:.2f} до {after['rps_k6']:.2f}, ошибок не появилось. Моя гипотеза: часть задержки возникла до начала обработки запроса приложением, возможно в клиенте k6 или сетевом пути NodePort. Это согласуется с большой разницей между p95 k6 и метрикой приложения. Чтобы доказать причину, нужны повторы теста и отдельное измерение ожидания на сетевом пути.")
paragraph('Для анализа я использовал запросы RPS, p95, CPU, памяти и ошибок из таблицы 2. Дополнительно p99 Grafana вычислялся заменой 0.95 на 0.99 в histogram_quantile. В сравнительной таблице выше p99 взят из k6, чтобы не смешивать место измерения.')

heading('5. Взаимосвязь метрик')
last = by_vu[200]
paragraph(f"При увеличении нагрузки с 10 до 200 VU показатель RPS k6 увеличился с {by_vu[10]['rps_k6']:.2f} до {last['rps_k6']:.2f} запросов в секунду. Одновременно CPU изменился с {by_vu[10]['cpu_cores']:.3f} до {last['cpu_cores']:.3f} ядра, память с {by_vu[10]['memory_mib']:.1f} до {last['memory_mib']:.1f} МиБ, а p95 k6 с {by_vu[10]['p95_k6_ms']:.2f} до {last['p95_k6_ms']:.2f} мс. Доля ошибок осталась 0%. Это показывает, что система продолжала обслуживать растущий поток, но клиент ждал ответ дольше. Одних этих графиков недостаточно, чтобы назвать конкретное узкое место.")

heading('6. Оповещение и экспорт dashboard')
paragraph('Я настроил правило Error Rate > 1% по запросу доли 5xx. Условие: значение выше 1%, период оценки 1 минута, время ожидания 2 минуты. Во время теста правило было в состоянии Normal, так как 5xx не наблюдались. Это проверка настройки, а не подтверждение срабатывания.')
picture('error-alert.jpg', 'Созданное правило Error Rate > 1% в Grafana')
paragraph('Я экспортировал модель dashboard в high-load-dashboard.json. Для переносимости из экспорта удалён только внутренний ID экземпляра Grafana. Через Dashboards → Import → Upload dashboard JSON file я загрузил этот файл как отдельную копию High-load Application (import check). После импорта Grafana открыла dashboard с семью панелями и теми же переменными.')
picture('dashboard-import-success.jpg', 'Успешно открытая копия dashboard после импорта JSON')

heading('7. Ответы на контрольные вопросы')
answers = [
    ('Что такое Grafana и для чего она используется?', 'Это интерфейс для просмотра и анализа метрик во времени, объединения графиков в dashboard и настройки оповещений.'),
    ('Чем Grafana отличается от Prometheus?', 'Prometheus регулярно собирает и хранит временные ряды, а Grafana запрашивает их и показывает в удобном виде.'),
    ('Какую роль выполняет Prometheus в системе мониторинга?', 'Он опрашивает /metrics приложения и компонентов Kubernetes, хранит точки рядов и отвечает на запросы PromQL.'),
    ('Что такое Dashboard и Panel?', 'Dashboard объединяет несколько визуализаций на одном экране. Panel показывает одну метрику или запрос.'),
    ('Для чего используется Time series?', 'Чтобы видеть, как значение меняется во времени, например рост RPS или задержки на этапах теста.'),
    ('Почему важно смотреть метрики во времени?', 'Среднее за весь тест скрывает момент изменения. Временная шкала позволяет сопоставить рост нагрузки с CPU, задержкой и ошибками.'),
    ('Как подключить Prometheus к Grafana?', 'В Connections → Data sources выбрать Prometheus, указать его URL и выполнить Test. В этой работе источник добавился вместе с Helm chart.'),
    ('Что такое PromQL и где он используется?', 'Это язык запросов Prometheus. В Grafana он используется в панелях и в правиле оповещения.'),
    ('Как получить RPS?', 'Взять скорость роста счётчика запросов: sum(rate(http_requests_total{namespace="vp-lr3",route="index"}[1m])).'),
    ('Как определить 5xx и их долю?', 'Скорость 5xx: sum(rate(http_requests_total{status=~"5.."}[1m])). Доля: 100 умножить эту скорость и разделить на скорость всех запросов.'),
    ('Почему одного CPU недостаточно для причины деградации?', 'CPU может расти вместе с полезной работой. Нужно проверить latency, RPS, ошибки, память и ограничения контейнера.'),
    ('Какие метрики сопоставлять при анализе?', 'RPS, p95 и p99 со стороны клиента и приложения, CPU, память, долю ошибок и состояние подов.'),
    ('Как увидеть приближение к saturation point?', 'При дальнейшем росте VU RPS перестаёт заметно увеличиваться, а задержка и часто ошибки продолжают расти.'),
    ('Почему RPS не обязан расти пропорционально нагрузке?', 'Есть ограничения CPU, очередей, сети и самого генератора нагрузки. Часть времени виртуальные пользователи ожидают ответ.'),
    ('Чем гипотеза об узком месте отличается от доказанного результата?', 'Рост p95 и CPU измерен. Причина задержки пока только предположение. Её надо проверить отдельными замерами и повторным тестом.'),
]
for number, (question, answer) in enumerate(answers, 1):
    paragraph(f'{number}. {question} {answer}')

heading('Вывод')
paragraph('Я подключил Grafana к Prometheus, создал и повторно импортировал dashboard, настроил оповещение и сопоставил пять этапов нагрузки. На 100 VU резко вырос клиентский p95, при этом RPS продолжил расти и ошибок не было даже на 200 VU. Эксперимент показывает ухудшение времени ответа для клиента, но не доказывает насыщение auth-service. Следующий шаг для проверки причины: повторить опыт и измерить ожидание до входа запроса в приложение.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Cm(0)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

doc.save(folder / 'ЛР3_отчёт.docx')
print('Saved', folder / 'ЛР3_отчёт.docx', figure_number, 'figures', table_number, 'tables')
