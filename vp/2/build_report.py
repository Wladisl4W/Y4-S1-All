import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


folder = Path(__file__).parent
data = json.loads((folder / 'results' / 'summary.json').read_text())
rows = data['rows']
comparison = data['comparison']
doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(3)
section.right_margin = Cm(1.5)
section.top_margin = Cm(2)
section.bottom_margin = Cm(1.3)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(14)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.space_after = Pt(5)

for style_name in ('Title', 'Heading 1'):
    style = doc.styles[style_name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.font.size = Pt(16 if style_name == 'Title' else 14)
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(5)


def paragraph(text):
    return doc.add_paragraph(text)


def heading(text):
    doc.add_paragraph(text, style='Heading 1')


def caption(text, above=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_before = Pt(6 if above else 2)
    p.paragraph_format.space_after = Pt(2 if above else 8)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    return p


def picture(name, title, width=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(folder / 'results' / name), width=Cm(width))
    caption(title)


def table(headers, items, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    t.autofit = False
    for i, head in enumerate(headers):
        t.rows[0].cells[i].text = head
    for item in items:
        cells = t.add_row().cells
        for i, value in enumerate(item):
            cells[i].text = str(value)
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cell.width = Cm(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.space_after = Pt(1)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10)
    return t


def command(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = 'Liberation Mono'
    run.font.size = Pt(9)


title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Лабораторная работа 2. Мониторинг приложения через Prometheus')
title_style_ppr = doc.styles['Title']._element.get_or_add_pPr()
border = title_style_ppr.find(qn('w:pBdr'))
if border is not None:
    title_style_ppr.remove(border)
title_ppr = title._p.get_or_add_pPr()
border = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'nil')
border.append(bottom)
title_ppr.append(border)
paragraph('Цель работы: настроить сбор метрик приложения, провести нагрузочные тесты и сравнить данные k6 с показателями Prometheus.')

heading('1. Объект и настройка мониторинга')
paragraph('Я тестировал страницу auth-service из проекта прошлого семестра. Приложение работает в локальном Kubernetes (K3s). Для этой работы я сделал отдельную копию Helm chart в папке ЛР2. Исходный проект: https://github.com/Wladisl4W/Y3-S2-Kubernetes.')
paragraph('В auth-service добавлены счётчик HTTP-запросов, гистограмма времени обработки и маршрут /metrics. Prometheus установлен в namespace monitoring через Helm chart kube-prometheus-stack. ServiceMonitor находит сервис auth-service и опрашивает /metrics каждые 5 секунд.')
command('helm upgrade --install prometheus ./vendor/kube-prometheus-stack-89.2.0.tgz -n monitoring --create-namespace -f prometheus-values.yaml')
command('KUBECONFIG=$HOME/.kube/config k3s kubectl apply -f service-monitor.yaml')
command('KUBECONFIG=$HOME/.kube/config k3s kubectl -n monitoring get pods')
picture('screenshots/monitoring-pods.png', 'Рисунок 1. Работающие поды мониторинга')
picture('screenshots/prometheus-targets.png', 'Рисунок 2. Target auth-service имеет состояние UP')
paragraph('В веб-интерфейсе запрос up{namespace="vp-lr2",service="auth-service"} вернул 1. Это значит, что Prometheus смог получить метрики сервиса во время последней проверки.')
picture('screenshots/prometheus-up.png', 'Рисунок 3. Результат запроса up для auth-service')

heading('2. Найденные метрики')
paragraph('В Prometheus я проверил метрики приложения, контейнера и подов. У всех строк в таблице есть реальные значения в сохранённом снимке metrics_snapshot.json. Ошибки 4xx и 5xx равны нулю, но оба счётчика созданы и доступны.')
caption('Таблица 1. Метрики, доступные в Prometheus', above=True)
table(['Метрика', 'Что измеряет', 'Тип', 'Единица'], [
    ('container_cpu_usage_seconds_total', 'Использованное контейнером процессорное время', 'Counter', 'сек'),
    ('process_cpu_seconds_total', 'Процессорное время процесса', 'Counter', 'сек'),
    ('container_memory_working_set_bytes', 'Рабочая память контейнера', 'Gauge', 'байт'),
    ('process_resident_memory_bytes', 'Физическая память процесса', 'Gauge', 'байт'),
    ('http_requests_total', 'Количество HTTP-запросов', 'Counter', 'запрос'),
    ('http_request_duration_seconds_bucket', 'Распределение времени обработки', 'Histogram', 'сек'),
    ('http_errors_total', 'Ошибки по группам 4xx и 5xx', 'Counter', 'ошибка'),
    ('kube_pod_status_phase', 'Фаза Pod: 0 или 1', 'Gauge', 'без единицы'),
    ('kube_pod_container_status_restarts_total', 'Число перезапусков контейнера', 'Counter', 'перезапуск'),
], [5.2, 6.1, 2.0, 1.7])

heading('3. Нагрузочные тесты')
paragraph('Я повторил сценарий ЛР1: GET-запрос к главной странице с проверкой статуса 200. Было три прогона по 30 секунд: 10, 50 и 200 виртуальных пользователей (VU). Между прогонами была пауза 40 секунд. Новые выводы k6 и ответы Prometheus сохранены в папке results.')
caption('Таблица 2. Результаты трёх новых прогонов', above=True)
table(['VU', 'RPS', 'p95, мс', 'p99, мс', 'Ошибки, %', 'CPU, %', 'Память, МиБ'], [
    (str(r['vus']), f"{r['rps']:.2f}", f"{r['p95_k6_ms']:.2f}", f"{r['p99_k6_ms']:.2f}",
     f"{r['errors_percent']:.0f}", f"{r['cpu_percent_of_one_core']:.1f}", f"{r['memory_mib']:.1f}") for r in rows
], [1.2, 2.0, 2.2, 2.2, 2.2, 2.0, 3.2])
paragraph('RPS, p95, p99 и доля ошибок взяты из k6. CPU и память взяты из Prometheus, усреднены по точкам после первых 15 секунд каждого прогона. CPU указан в процентах от одного логического ядра. Поэтому 115,4 % означает примерно 1,15 ядра, а не невозможное значение.')
paragraph('Основные запросы PromQL, которые я использовал:')
command('100 * sum(rate(container_cpu_usage_seconds_total{namespace="vp-lr2",container="auth-service"}[20s]))')
command('sum(container_memory_working_set_bytes{namespace="vp-lr2",container="auth-service"}) / 1024 / 1024')
command('sum(rate(http_requests_total{namespace="vp-lr2",route="index"}[15s]))')
command('sum(rate(http_errors_total{namespace="vp-lr2",group="5xx"}[15s])) or vector(0)')
command('1000 * histogram_quantile(0.95, sum by (le) (increase(http_request_duration_seconds_bucket{namespace="vp-lr2",route="index"}[40s])))')
paragraph('Ниже показаны настоящие снимки терминала. На них открыт сохранённый вывод именно этих трёх прогонов, а не запуск теста в момент снимка.')
for number, vus in enumerate((10, 50, 200), 4):
    picture(f'screenshots/vu{vus}.png', f'Рисунок {number}. Сохранённый вывод k6 для {vus} VU')

heading('4. Сравнение и анализ')
paragraph(f"Для 200 VU k6 показал p95 = {rows[2]['p95_k6_ms']:.2f} мс, а Prometheus = {rows[2]['p95_prometheus_ms']:.2f} мс. Разница равна {comparison['difference_ms']:.2f} мс, или {comparison['difference_percent_of_k6']:.2f} % от значения k6. k6 измеряет полный ответ клиенту. Метрика приложения измеряет только обработку внутри Flask. Кроме того, Prometheus оценивает квантиль по бакетам за окно 40 секунд, а k6 считает по ответам во время 30-секундного теста. Поэтому значения напрямую не совпадают.")
picture('p95.png', 'Рисунок 7. Изменение p95 при росте нагрузки')
picture('resources.png', 'Рисунок 8. Средние CPU и память auth-service')
picture('rps.png', 'Рисунок 9. Пропускная способность по данным k6')
paragraph('Ухудшение заметно на 200 VU: p95 k6 вырос с 49,96 до 228,16 мс, а RPS вырос с 48,56 до 176,09, то есть не пропорционально числу VU. CPU вырос с 33,3 до 115,4 % одного ядра. Память изменилась мало: с 31,4 до 33,5 МиБ. Ошибок не было. Ранее, уже на 50 VU, p95 вырос с 25,07 до 49,96 мс, но более резкий рост виден между 50 и 200 VU.')
paragraph('Моя гипотеза: при высокой нагрузке запросы дольше ждут до начала обработки во Flask или в сетевом стеке. На это указывает большая разница между p95 k6 и p95 приложения. Рост CPU тоже может участвовать, но по этим данным нельзя доказать, что процессор исчерпан: контейнер не ограничен одним ядром. Для проверки гипотезы нужны дополнительные измерения очереди, профилирование и повторные прогоны.')

heading('5. Ответы на контрольные вопросы')
answers = [
    ('Что такое мониторинг и чем он отличается от нагрузочного тестирования?', 'Мониторинг постоянно собирает состояние работающей системы. Нагрузочное тестирование специально создаёт запросы и показывает, как система ведёт себя при заданной нагрузке.'),
    ('Что дают k6 и Prometheus и почему RPS и p95 могут различаться?', 'k6 видит систему как клиент: запросы, ошибки и полное время ответа. Prometheus собирает внутренние метрики. Различия возникают из-за места измерения, интервалов и способа вычисления квантилей.'),
    ('Что такое Prometheus и Pull-модель?', 'Prometheus хранит временные ряды и сам регулярно забирает метрики с HTTP-адреса /metrics. Это и есть Pull-модель.'),
    ('Что такое метрика, временной ряд и labels?', 'Метрика описывает измеряемую величину. Временной ряд состоит из её значений с отметками времени. Labels уточняют, к какому сервису, поду или маршруту относится ряд.'),
    ('Чем counter отличается от gauge и почему rate() не применяют к gauge?', 'Counter только растёт до перезапуска, например число запросов. Gauge может расти и падать, например память. rate() нужен для скорости роста счётчика и обрабатывает его сброс; для обычного gauge такая интерпретация неверна.'),
    ('Почему для latency используют Histogram? Чем он отличается от Summary?', 'Histogram считает наблюдения в бакетах, которые можно складывать между экземплярами и затем оценивать p95 через histogram_quantile(0.95, sum by (le) (rate(http_request_duration_seconds_bucket[5m]))). Summary вычисляет квантили на стороне приложения; такие квантили между экземплярами корректно не складываются.'),
    ('Что такое PromQL и как найти 5xx-ошибки и их долю?', 'PromQL это язык запросов Prometheus. Скорость 5xx: sum(rate(http_requests_total{status=~"5.."}[1m])). Доля в процентах: 100 * эта скорость / sum(rate(http_requests_total[1m])). Для количества за минуту можно использовать sum(increase(http_requests_total{status=~"5.."}[1m])).'),
    ('Что такое bottleneck и почему CPU 90 % не доказывает его?', 'Bottleneck это ресурс или участок, ограничивающий скорость работы. Одно значение CPU не показывает, мешает ли именно процессор: нужны задержки, пропускная способность, ограничения контейнера и другие метрики.'),
    ('Что такое warm-up и saturation point?', 'Warm-up это первые секунды, когда приложение и инструменты ещё прогреваются. Saturation point это нагрузка, после которой рост запросов почти не увеличивает полезную производительность, а задержка растёт. Первые секунды лучше отдельно не усреднять с устойчивой работой.'),
    ('Чем гипотеза отличается от доказанного результата?', 'Измеренный p95 и CPU это факты. Причина их изменения пока только объяснение, которое надо проверить отдельным экспериментом. Я поэтому пишу «возможно, запросы ждут до обработки», а не утверждаю, что это доказано.'),
]
for i, (question, answer) in enumerate(answers, 1):
    paragraph(f'{i}. {question} {answer}')

heading('Вывод')
paragraph('Я настроил Prometheus для auth-service и получил метрики приложения и Kubernetes. Все три теста прошли без ошибок. При 200 VU задержка со стороны клиента выросла намного сильнее, чем измеренное внутри приложения время обработки. Это даёт направление для дальнейшей проверки, но пока не доказывает причину замедления.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Cm(0)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

doc.save(folder / 'ЛР2_отчёт.docx')
