import csv
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = Path(__file__).parent
rows = list(csv.DictReader((HERE / 'critical_delay.csv').open()))
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 11})
fig, ax = plt.subplots(figsize=(8, 4.5))
for a1 in ['0.5', '1.0', '1.5']:
    series = [r for r in rows if float(r['a1']) == float(a1)]
    ax.plot([float(r['k']) for r in series], [float(r['critical_tau']) for r in series], marker='o', label='a1 = ' + a1)
ax.set(xlabel='Усиление регулятора k', ylabel='Критическая задержка τ, с')
ax.grid(alpha=.3); ax.legend(); fig.tight_layout()
fig.savefig(HERE / 'critical_delay.png', dpi=180); plt.close(fig)

fig, ax = plt.subplots(figsize=(8, 4.5))
for file, label in [('response_no_delay.csv', 'τ = 0'), ('response_stable.csv', 'τ = 0,415 с'), ('response_unstable.csv', 'τ = 0,715 с')]:
    data = list(csv.DictReader((HERE / file).open()))
    points = [r for r in data if float(r['time']) <= 35]
    ax.plot([float(r['time']) for r in points], [float(r['x']) for r in points], label=label)
ax.axhline(2/3, color='black', linestyle='--', linewidth=1, label='Равновесие x = 2/3')
ax.set(xlabel='Время t, с', ylabel='Состояние x(t)', ylim=(-5, 5))
ax.grid(alpha=.3); ax.legend(); fig.tight_layout()
fig.savefig(HERE / 'responses.png', dpi=180); plt.close(fig)

doc = Document()
s = doc.sections[0]
s.page_width = Cm(21); s.page_height = Cm(29.7)
s.left_margin = Cm(3); s.right_margin = Cm(1.5)
s.top_margin = Cm(2); s.bottom_margin = Cm(1.5)
style = doc.styles['Normal']; style.font.name = 'Times New Roman'; style.font.size = Pt(14)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.first_line_indent = Cm(1.25)
style.paragraph_format.space_after = Pt(4)
for name in ['Heading 1']:
    st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.size = Pt(14); st.font.bold = True; st.font.color.rgb = None
    st.paragraph_format.space_before = Pt(10); st.paragraph_format.space_after = Pt(5)

def para(text): doc.add_paragraph(text)
def caption(text):
    p=doc.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    for run in p.runs: run.font.size=Pt(12)
def figure(file, title):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(HERE / file), width=Cm(13.5))
    caption(title)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent=Cm(0)
r=p.add_run('Лабораторная работа. Устойчивость динамической системы с агентом')
r.font.name='Times New Roman'; r.font.size=Pt(16); r.font.bold=True
para('Цель работы: смоделировать систему второго порядка с задержкой измерения и определить, при какой задержке колебания перестают затухать.')
doc.add_heading('1. Модель системы', 1)
para('Объект описывается уравнением a2·x″ + a1·x′ + a0·x = k·(g − x(t−τ)). Здесь x(t) — состояние системы, g(t) — заданная цель, k — усиление регулятора, τ — задержка измерения. Я принял g(t)=1 при t≥0 и нулевое начальное состояние. В основном опыте a2=1, a0=1, а a1 и k меняются.')
para('Из-за задержки регулятор сравнивает цель с прошлым состоянием. При слишком большой задержке он исправляет уже устаревшую ошибку, поэтому колебания могут расти.')
doc.add_heading('2. Метод расчёта', 1)
para('В программе время разбито на шаги Δt=0,01 с. На каждом шаге я беру сохранённое значение x(t−τ), вычисляю ускорение по уравнению, затем обновляю скорость и состояние. Для границы устойчивости использован двоичный поиск по τ.')
para('Сравниваются средние отклонения от равновесия в интервалах 180–230 с и 250–300 с. Если позднее отклонение больше раннего, колебания считаются растущими. Граница уточняется до ширины интервала 0,01 с. Это численная оценка для выбранного шага и времени наблюдения.')
doc.add_heading('3. Результаты', 1)
caption('Таблица 1. Оценка критической задержки τ, с')
table=doc.add_table(rows=1, cols=4); table.style='Table Grid'; table.alignment=WD_TABLE_ALIGNMENT.CENTER
for cell,text in zip(table.rows[0].cells,['k','a1 = 0,5','a1 = 1,0','a1 = 1,5']): cell.text=text
for k in ['1.5','2.0','2.5','3.0','3.5','4.0']:
    cells=table.add_row().cells; cells[0].text=k.replace('.',',')
    for j,a1 in enumerate(['0.5','1.0','1.5'],1):
        val=next(float(r['critical_tau']) for r in rows if float(r['k'])==float(k) and float(r['a1'])==float(a1))
        cells[j].text=f'{val:.3f}'.replace('.',',')
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Cm(0)
            for run in p.runs: run.font.size=Pt(11)
figure('critical_delay.png','Рисунок 1. Критическая задержка при разных a1 и k')
para('При росте k допустимая задержка уменьшается. При большем a1 колебания затухают сильнее, поэтому система выдерживает большую задержку.')
figure('responses.png','Рисунок 2. Изменение состояния при k=2 и a1=1')
para('Для k=2 и a1=1 граница составляет примерно 0,565 с. При τ=0,415 с колебания затухают и x стремится к 2/3. При τ=0,715 с их размах растёт. Цель равна 1, но равновесное x=2/3, так как в модели присутствует член a0·x.')
doc.add_heading('4. Как работает программа', 1)
para('Сначала программа задаёт коэффициенты модели и шаг времени. Функция simulate хранит историю x в массиве, чтобы взять значение с задержкой. В цикле она вычисляет ускорение, изменяет скорость и новое состояние. Затем сравнивает размах колебаний в двух поздних промежутках времени.')
para('Функция criticalDelay подбирает задержку: сначала находит устойчивый и неустойчивый варианты, затем делит интервал пополам. Главная функция повторяет этот расчёт для всех пар a1 и k, записывает таблицу CSV и три временных ряда. Отчёт строится по этим сохранённым файлам.')
doc.add_heading('5. Вывод', 1)
para('Я смоделировал работу регулятора с запаздывающим датчиком. При увеличении задержки устойчивость теряется. Для выбранных коэффициентов усиление регулятора уменьшает допустимую задержку, а увеличение a1 повышает её. Точные числа зависят от параметров модели и шага расчёта.')
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent=Cm(0)
doc.save(HERE / 'ЛР6_отчёт.docx')
