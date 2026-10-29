import csv
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

folder = Path(__file__).parent
rows = list(csv.DictReader((folder / 'free_probability.csv').open()))
plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 11})
fig, ax = plt.subplots(figsize=(8, 4.4))
for case in range(1, 6):
 series = [r for r in rows if int(r['case']) == case]
 label = f"{case}: λ={float(series[0]['lambda']):g}, μ={float(series[0]['mu']):g}"
 ax.plot([int(r['time']) for r in series], [float(r['free_probability']) for r in series], label=label)
ax.set(xlabel='Момент времени t', ylabel='Доля запусков со свободным сервером', xlim=(0,30), ylim=(0,1.05))
ax.grid(alpha=.3); ax.legend(fontsize=8); fig.tight_layout()
fig.savefig(folder / 'free_probability.png', dpi=180); plt.close(fig)

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.left_margin=Cm(3); sec.right_margin=Cm(1.5)
sec.top_margin=Cm(2); sec.bottom_margin=Cm(1.5)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(14)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.first_line_indent=Cm(1.25)
normal.paragraph_format.space_after=Pt(4)
head=doc.styles['Heading 1']; head.font.name='Times New Roman'; head.font.size=Pt(14); head.font.bold=True; head.font.color.rgb=RGBColor(0,0,0)
head.paragraph_format.space_before=Pt(9); head.paragraph_format.space_after=Pt(4)

def para(text): return doc.add_paragraph(text)
def cap(text):
 p=doc.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
 for r in p.runs: r.font.size=Pt(12)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
r=p.add_run('Лабораторная работа. Вероятность свободного сервера в системе M/M/1')
r.font.name='Times New Roman'; r.font.size=Pt(16); r.font.bold=True
para('Цель работы: оценить моделированием вероятность того, что один сервер свободен в заданный момент времени, и сравнить результат с теоретическим пределом.')
doc.add_heading('1. Модель и исходные данные',1)
para('Я рассматриваю систему M/M/1: один сервер и очередь без ограничения длины. Интервалы между поступлениями и длительность обслуживания имеют экспоненциальное распределение. Для сравнения я провёл пять серий с разными интенсивностями поступления λ и обслуживания μ. В начале каждого запуска система пуста.')
para('Во всех пяти случаях λ<μ, поэтому существует стационарный режим. Теоретическая доля времени со свободным сервером равна 1−λ/μ. Для каждого набора коэффициентов она своя.')
doc.add_heading('2. Ход моделирования',1)
para('Для каждого набора коэффициентов и каждого целого t от 0 до 30 я провёл 10 000 независимых запусков с пустой системы. Модель обрабатывала поступления и окончания обслуживания до момента t, после чего я проверял, свободен ли сервер.')
para('Оценка вероятности равна числу запусков со свободным сервером, делённому на 10 000. Например, для случая 2 при t=10 получилось 4343 свободных запуска, то есть 0,4343. Полная таблица сохранена в free_probability.csv.')
h=doc.add_heading('3. Результаты',1); h.paragraph_format.page_break_before=True
cap('Таблица 1. Сравнение пяти наборов коэффициентов')
tab=doc.add_table(rows=1, cols=6); tab.style='Table Grid'; tab.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,s in zip(tab.rows[0].cells,['Случай','λ','μ','Предел','t=10','t=30']): c.text=s
for case in range(1, 6):
 series=[r for r in rows if int(r['case'])==case]
 first=series[0]
 at10=next(r for r in series if int(r['time'])==10)
 at30=next(r for r in series if int(r['time'])==30)
 cells=tab.add_row().cells
 values=[str(case),first['lambda'],first['mu'],f"{1-float(first['lambda'])/float(first['mu']):.4f}",at10['free_probability'],at30['free_probability']]
 for c,s in zip(cells,values): c.text=s.replace('.',',')
for row in tab.rows:
 for cell in row.cells:
  for p in cell.paragraphs:
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0)
   for run in p.runs: run.font.size=Pt(11)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent=Cm(0); p.paragraph_format.keep_with_next=True
p.add_run().add_picture(str(folder / 'free_probability.png'),width=Cm(12.5))
cap('Рисунок 1. Вероятность свободного сервера в зависимости от времени')
para('В момент t=0 сервер всегда свободен. Затем во всех сериях доля свободных состояний уменьшается и приближается к своему пределу. Чем больше отношение λ/μ, тем реже сервер свободен. Небольшие колебания кривых объясняются случайностью и конечным числом повторений.')
doc.add_heading('4. Вывод',1)
para('Я оценил вероятность свободного сервера для пяти наборов коэффициентов. Результаты при больших t близки к соответствующим теоретическим значениям 1−λ/μ. Для каждого t опыты начинались заново с пустой системы.')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.paragraph_format.first_line_indent=Cm(0)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.save(folder / 'ЛР7_отчёт.docx')
